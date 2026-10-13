"""元大薪資檔預覽、下載及上傳待放行流程。"""

from __future__ import annotations

import io
import re
import tempfile
import time
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import load_workbook

from modules.auth import get_jenny_drive_service


INTERNAL_SALARY_ROOT_ID = "1LYmhOIcgjmO5-74zh1D1pKfVizIF5Wls"
FOLDER_MIME = "application/vnd.google-apps.folder"
YUANTA_URL = "https://b2bank.yuantabank.com.tw/B2C/login/LOGIN_Home.faces"
FUBON_URL = "https://ebank.taipeifubon.com.tw/B2C/common/Index.faces"
SUPPORTED_REGIONS = ("台北", "台中")


@dataclass(frozen=True)
class SalaryFile:
    file_id: str
    name: str
    amount: int
    kind: str
    summary: str


def _emit(log: Callable[[str], None] | None, message: str) -> None:
    if log:
        log(message)


def _find_child(drive, parent_id: str, name: str, mime_type: str | None = None):
    escaped_name = name.replace("\\", "\\\\").replace("'", "\\'")
    query = f"'{parent_id}' in parents and name='{escaped_name}' and trashed=false"
    if mime_type:
        query += f" and mimeType='{mime_type}'"
    rows = drive.files().list(
        q=query,
        fields="files(id,name,mimeType,modifiedTime)",
        supportsAllDrives=True,
        includeItemsFromAllDrives=True,
        orderBy="modifiedTime desc",
        pageSize=20,
    ).execute().get("files", [])
    return rows[0] if rows else None


def _download(drive, file_id: str) -> bytes:
    return drive.files().get_media(fileId=file_id).execute()


def _number(value) -> Decimal:
    text = str(value or "").strip().replace(",", "").replace("$", "")
    if not text:
        return Decimal(0)
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal(0)


def xlsx_total(data: bytes) -> int:
    """加總元大匯入檔 D3:D。"""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    worksheet = workbook["元大"] if "元大" in workbook.sheetnames else workbook.active
    total = sum(
        (_number(row[0]) for row in worksheet.iter_rows(
            min_row=3, min_col=4, max_col=4, values_only=True
        )),
        Decimal(0),
    )
    if total != total.to_integral_value():
        raise ValueError(f"薪資檔金額不是整數：{total}")
    return int(total)


def _salary_file(drive, item: dict, kind: str, summary: str) -> SalaryFile:
    return SalaryFile(
        item["id"], item["name"], xlsx_total(_download(drive, item["id"])), kind, summary
    )


def build_salary_plan(
    root_folder_id: str,
    period: str,
    region: str,
    include_internal_salary: bool = False,
) -> list[SalaryFile]:
    if region not in SUPPORTED_REGIONS:
        raise ValueError("元大薪資檔上傳目前只支援台北、台中")
    if not re.fullmatch(r"\d{6}-[12]", period):
        raise ValueError("期別格式須為 YYYYMM-1 或 YYYYMM-2")

    drive = get_jenny_drive_service()
    period_folder = _find_child(drive, root_folder_id, period, FOLDER_MIME)
    if not period_folder:
        raise FileNotFoundError(f"找不到期別資料夾：{period}")

    fee_name = f"{period}元大承攬費-{region}.xlsx"
    fee = _find_child(drive, period_folder["id"], fee_name)
    if not fee:
        raise FileNotFoundError(f"找不到必要檔案：{fee_name}")
    files = [_salary_file(drive, fee, "承攬費", "員工酬勞")]

    # 工具包押金不需要勾選：該期別有檔案就自動加入，沒有就略過。
    deposit_name = f"{period}元大工具包押金-{region}.xlsx"
    deposit = _find_child(drive, period_folder["id"], deposit_name)
    if deposit:
        files.append(_salary_file(drive, deposit, "工具包押金", "費用"))

    if include_internal_salary:
        yyyymm = period[:6]
        year_name = f"{yyyymm[:4]}薪資"
        year_folder = _find_child(drive, INTERNAL_SALARY_ROOT_ID, year_name, FOLDER_MIME)
        if not year_folder:
            raise FileNotFoundError(f"找不到內勤薪資資料夾：{year_name}")
        region_folder = _find_child(drive, year_folder["id"], region, FOLDER_MIME)
        if not region_folder:
            raise FileNotFoundError(f"找不到內勤薪資地區資料夾：{region}")
        internal_name = f"{yyyymm}元大帳戶-{region}.xlsx"
        internal = _find_child(drive, region_folder["id"], internal_name)
        if not internal:
            raise FileNotFoundError(f"找不到內勤薪資檔：{internal_name}")
        files.append(_salary_file(drive, internal, "內勤薪資", "薪資"))
    return files


def plan_total(files: Iterable[SalaryFile]) -> int:
    return sum(item.amount for item in files)


def _all_frames(page):
    return [frame for candidate in page.context.pages if not candidate.is_closed()
            for frame in candidate.frames]


def _visible(locator):
    try:
        for index in range(locator.count()):
            item = locator.nth(index)
            if item.is_visible():
                return item
    except Exception:
        pass
    return None


def _click_text(page, text: str, timeout: int = 30) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        for frame in _all_frames(page):
            for exact in (True, False):
                item = _visible(frame.get_by_text(text, exact=exact))
                if item is not None:
                    item.click()
                    return
        page.wait_for_timeout(500)
    raise RuntimeError(f"銀行頁面找不到「{text}」")


def _row_for_label(page, label: str, timeout: int = 30):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        for frame in _all_frames(page):
            node = _visible(frame.get_by_text(label, exact=True))
            if node is None:
                continue
            for xpath in ("ancestor::tr[1]", "ancestor::div[contains(@class,'row')][1]", "parent::*"):
                row = node.locator(f"xpath={xpath}")
                if row.count():
                    return row.first
        page.wait_for_timeout(300)
    raise RuntimeError(f"銀行頁面找不到欄位「{label}」")


def _select_option(page, text: str, label_hint: str = "") -> None:
    for frame in _all_frames(page):
        selects = frame.locator("select")
        for index in range(selects.count()):
            select = selects.nth(index)
            if not select.is_visible():
                continue
            for option_index in range(select.locator("option").count()):
                option = select.locator("option").nth(option_index)
                if text not in (option.inner_text() or ""):
                    continue
                if label_hint:
                    row = select.locator("xpath=ancestor::*[self::tr or self::div][1]")
                    if row.count() and label_hint not in row.inner_text():
                        continue
                select.select_option(value=option.get_attribute("value"))
                return
    raise RuntimeError(f"銀行頁面找不到選項「{text}」")


def _choose_account(page, row_label: str, target: str) -> None:
    row = _row_for_label(page, row_label)
    if target in row.inner_text() and "請選擇" not in row.inner_text():
        return
    trigger = _visible(row.get_by_text(re.compile("請選擇"), exact=False))
    if trigger is None:
        raise RuntimeError(f"{row_label}找不到帳號選擇按鈕")
    trigger.click(force=True)
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        for frame in _all_frames(page):
            card = _visible(frame.get_by_text(target, exact=False))
            if card is not None:
                card.click(force=True)
                page.wait_for_timeout(1000)
                if "請選擇" not in _row_for_label(page, row_label).inner_text():
                    return
        page.wait_for_timeout(300)
    raise RuntimeError(f"{row_label}找不到常用帳號「{target}」")


def _connect(cdp_url: str):
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError("缺少 playwright；請安裝套件並重新啟動薪資系統") from exc
    runtime = sync_playwright().start()
    try:
        browser = runtime.chromium.connect_over_cdp(cdp_url)
    except Exception as exc:
        runtime.stop()
        raise RuntimeError("無法連線本機 Chrome（remote-debugging-port=9222）") from exc
    return runtime, browser


def _page(browser, host: str, url: str):
    context = browser.contexts[0] if browser.contexts else browser.new_context()
    for page in context.pages:
        if host in page.url:
            page.bring_to_front()
            return page
    page = context.new_page()
    page.goto(url, wait_until="domcontentloaded", timeout=30_000)
    return page


def _wait_login(page, marker: str, bank: str, timeout: int = 300) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        for frame in _all_frames(page):
            try:
                if marker in (frame.locator("body").inner_text(timeout=1000) or ""):
                    return
            except Exception:
                pass
        page.wait_for_timeout(1000)
    raise RuntimeError(f"等待{bank}登入逾時")


def upload_yuanta_salary_files(
    files: list[SalaryFile], region: str,
    cdp_url: str = "http://127.0.0.1:9222",
    log: Callable[[str], None] | None = None,
) -> None:
    if not files:
        raise ValueError("沒有可上傳的薪資檔")
    drive = get_jenny_drive_service()
    with tempfile.TemporaryDirectory(prefix="yuanta-salary-") as tmp:
        staged = []
        for item in files:
            path = Path(tmp) / item.name
            path.write_bytes(_download(drive, item.file_id))
            staged.append((item, path))
        runtime, browser = _connect(cdp_url)
        try:
            page = _page(browser, "b2bank.yuantabank.com.tw", YUANTA_URL)
            _emit(log, "請在 Chrome 完成元大登入；登入後會自動繼續")
            _wait_login(page, "收付款服務", "元大")
            for item, path in staged:
                _emit(log, f"上傳中：{item.name}（NT$ {item.amount:,}）")
                _click_text(page, "收付款服務")
                _click_text(page, "薪資付款")
                _click_text(page, "薪資臺幣整批付款")
                _select_option(page, "臺幣薪轉格式Excel")
                uploaded = False
                for frame in _all_frames(page):
                    field = _visible(frame.locator("input[type=file]"))
                    if field is not None:
                        field.set_input_files(str(path))
                        uploaded = True
                        break
                if not uploaded:
                    raise RuntimeError("元大頁面找不到檔案上傳欄位")
                _click_text(page, "確定")
                _select_option(page, "元大大同綜合活期", "付款帳號")
                _select_option(page, item.summary, "摘要")
                _click_text(page, "審核")
                _click_text(page, "確認審核")
                _emit(log, f"已送待放行：{item.name}")
        finally:
            runtime.stop()


def prepare_fubon_to_yuanta(
    region: str, amount: int,
    cdp_url: str = "http://127.0.0.1:9222",
    log: Callable[[str], None] | None = None,
) -> None:
    if region not in SUPPORTED_REGIONS or amount <= 0:
        raise ValueError("富邦匯款到元大的地區或金額不正確")
    runtime, browser = _connect(cdp_url)
    try:
        page = _page(browser, "ebank.taipeifubon.com.tw", FUBON_URL)
        _emit(log, "請在 Chrome 完成富邦登入；登入後會自動繼續")
        _wait_login(page, "臺幣轉帳", "富邦")
        _click_text(page, "臺幣轉帳")
        _click_text(page, "立即/預約轉帳")
        if region == "台北":
            _choose_account(page, "轉出帳號", "松高分行")
        _choose_account(page, "轉入帳號", f"元大{region}")
        row = _row_for_label(page, "轉帳金額")
        field = _visible(row.locator(
            'input:not([type=hidden]):not([type=radio]):not([type=button]):not([type=submit])'
        ))
        if field is None:
            raise RuntimeError("富邦頁面找不到轉帳金額欄位")
        field.fill(str(amount))
        _emit(log, f"富邦已預填元大{region}與 NT$ {amount:,}；請人工核對並確認送出")
    finally:
        runtime.stop()
